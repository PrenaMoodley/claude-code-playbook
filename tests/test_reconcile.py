import importlib.util
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / ".claude/skills/reconcile-files/scripts/reconcile.py"
spec = importlib.util.spec_from_file_location("reconcile", SCRIPT)
reconcile = importlib.util.module_from_spec(spec)
spec.loader.exec_module(reconcile)


def write(path: Path, text: str) -> Path:
    path.write_text(text, encoding="utf-8")
    return path


def test_amount_parsing():
    assert reconcile.to_decimal("R 1,234.50") == Decimal("1234.50")
    assert reconcile.to_decimal("(99.99)") == Decimal("-99.99")
    assert reconcile.to_decimal("0.005") == Decimal("0.01")


def test_reconcile_end_to_end(tmp_path, capsys):
    bank = write(
        tmp_path / "bank.csv",
        "ref,amount\nINV-1,100.00\nINV-2,50.00\nINV-3,10.00\nINV-3,10.00\n",
    )
    ledger = write(
        tmp_path / "ledger.csv",
        "invoice,value\nINV-1,-100.00\nINV-2,-45.00\nINV-4,-30.00\n",
    )
    out = tmp_path / "rec.xlsx"
    reconcile.main(
        [str(bank), str(ledger), "--left-key", "ref", "--right-key", "invoice",
         "--left-amount", "amount", "--right-amount", "value", "--right-negate", "--out", str(out)]
    )  # fmt: skip

    printed = capsys.readouterr().out
    assert "duplicate keys" in printed

    wb = load_workbook(out)
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True)}
    assert summary["Matched keys"] == 1
    assert summary["Net difference"] == Decimal("-5.00")  # 170 - 175
    assert [r for r in wb["Differences"].iter_rows(min_row=2, values_only=True)] == [
        ("INV-2", Decimal("50.00"), Decimal("45.00"), Decimal("5.00"))
    ]
    assert [r[0] for r in wb["Only left"].iter_rows(min_row=2, values_only=True)] == ["INV-3"]
    assert [r[0] for r in wb["Only right"].iter_rows(min_row=2, values_only=True)] == ["INV-4"]
    assert [r[0] for r in wb["Duplicates"].iter_rows(min_row=2, values_only=True)] == ["INV-3"]
